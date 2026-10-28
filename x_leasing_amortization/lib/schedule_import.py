"""Contract-schedule input validation, without computing an interest formula."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


HEADER_COLUMNS = (
    'Import Key', 'Compute Method', 'Agreement No.', 'Leasing Contract No.',
    'Vendor', 'Plate Number', 'Start Date', 'First Due Date', 'Duration (Months)',
    'Total Hutang', 'Harga OTR', 'Down Payment', 'Monthly Installment',
    'Interest (%/Year)', 'Total Bunga', 'Advarr', 'Principal Account',
    'Short Term Account', 'Interest Expense Account', 'Accrued Interest Account', 'Interest Journal',
)
LINE_COLUMNS = ('Import Key', 'Installment No.', 'Due Date', 'Principal', 'Interest',
                'Total', 'Principal Balance', 'Interest Balance')
OPTIONAL_HEADER = {'Leasing Contract No.', 'Harga OTR', 'Down Payment'}
OPTIONAL_LINE = {'Total', 'Principal Balance', 'Interest Balance'}
METHODS = {
    'Bunga Efektif': 'effective', 'Bunga Efektif (Tanpa DP)': 'effective_2', 'Bunga Flat': 'flat',
    'Bunga Efektif 2': 'effective_2',  # Compatibility with previously downloaded templates.
}
ADVARR = {'ADDM': 'in_advance_addm', 'Arrears': 'in_arrears'}
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 12000


class ScheduleImportError(ValueError):
    pass


def number(value, label, *, quantum=None, integer=False, positive=False):
    if isinstance(value, bool) or value is None or value == '':
        raise ScheduleImportError(f'{label}: isi angka, termasuk 0 bila nilainya nol.')
    try:
        result = Decimal(str(value))
    except InvalidOperation as error:
        raise ScheduleImportError(f'{label}: gunakan angka tanpa Rp, pemisah ribuan atau persen.') from error
    if not result.is_finite() or result < 0 or (positive and result <= 0):
        raise ScheduleImportError(f'{label}: angka harus valid dan {"lebih besar dari nol" if positive else "tidak negatif"}.')
    if integer and result != result.to_integral_value():
        raise ScheduleImportError(f'{label}: harus berupa bilangan bulat.')
    try:
        if quantum is not None and result % Decimal(str(quantum)):
            raise ScheduleImportError(f'{label}: nominal harus sesuai presisi mata uang {quantum}; tidak dibulatkan otomatis.')
    except InvalidOperation as error:
        raise ScheduleImportError(f'{label}: nominal terlalu besar untuk diproses.') from error
    return result


def input_date(value, label):
    if isinstance(value, datetime):
        if value.time() != datetime.min.time():
            raise ScheduleImportError(f'{label}: tanggal tidak boleh memiliki jam.')
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(value) if isinstance(value, str) else _bad_date()
    except ValueError as error:
        raise ScheduleImportError(f'{label}: gunakan tanggal Excel atau YYYY-MM-DD.') from error


def _bad_date():
    raise ValueError('Missing date')


def read_workbook(raw):
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ScheduleImportError('Upload file XLSX, maksimal 5 MB.')
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if (len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024
                    or any(item.filename.endswith('vbaProject.bin') for item in archive.infolist())):
                raise ScheduleImportError('Ukuran isi file terlalu besar atau file mengandung macro.')
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, ValueError, ParseError) as error:
        if isinstance(error, ScheduleImportError):
            raise
        raise ScheduleImportError('File bukan workbook XLSX yang valid.') from error
    try:
        result = {}
        for name, columns, optional in [('Header', HEADER_COLUMNS, OPTIONAL_HEADER),
                                         ('Angsuran', LINE_COLUMNS, OPTIONAL_LINE)]:
            if name not in workbook.sheetnames:
                raise ScheduleImportError(f'Sheet {name} tidak ditemukan. Gunakan template resmi.')
            sheet = workbook[name]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(columns):
                raise ScheduleImportError(f'{name}: maksimal {MAX_ROWS} baris dan {len(columns)} kolom.')
            rows = sheet.iter_rows()
            header = [cell.value for cell in next(rows)]
            if len(set(header)) != len(header) or set(header) != set(columns):
                raise ScheduleImportError(f'{name}: nama kolom harus sama dengan template, tanpa duplikasi.')
            data = []
            for index, cells in enumerate(rows, 2):
                if index > MAX_ROWS + 1 or len(cells) > len(columns):
                    raise ScheduleImportError(f'{name}: terlalu banyak baris/kolom.')
                if all(cell.value is None or cell.value == '' for cell in cells):
                    continue
                if any(cell.data_type == 'f' for cell in cells):
                    raise ScheduleImportError(f'{name} baris {index}: formula tidak diterima. Paste Special > Values sebelum upload.')
                record = {label: cell.value.strip() if isinstance(cell.value, str) else cell.value
                          for label, cell in zip(header, cells)}
                for label in set(columns) - optional:
                    if record.get(label) is None or record.get(label) == '':
                        raise ScheduleImportError(f'{name} baris {index}: {label} wajib diisi.')
                record['_row'] = index
                data.append(record)
            if not data:
                raise ScheduleImportError(f'Sheet {name} belum berisi data.')
            result[name] = data
        if len(result['Header']) > 100:
            raise ScheduleImportError('Maksimal 100 kontrak per upload.')
        return result
    finally:
        workbook.close()


def validate_rows(header, rows, quantum):
    """Return canonical payment rows; accept the Flat first-row zero display."""
    label = f"Import Key {header['Import Key']}"
    duration = int(number(header['Duration (Months)'], f'{label} Tenor', integer=True, positive=True))
    principal = number(header['Total Hutang'], f'{label} Total Hutang', quantum=quantum, positive=True)
    total_interest = number(header['Total Bunga'], f'{label} Total Bunga', quantum=quantum)
    installment = number(header['Monthly Installment'], f'{label} Cicilan', quantum=quantum, positive=True)
    rate = number(header['Interest (%/Year)'], f'{label} Interest')
    due = input_date(header['First Due Date'], f'{label} First Due Date')
    start = input_date(header['Start Date'], f'{label} Start Date')
    if start > due:
        raise ScheduleImportError(f'{label}: Start Date tidak boleh setelah First Due Date.')
    if len(rows) != duration:
        raise ScheduleImportError(f'{label}: jumlah baris angsuran {len(rows)} tidak sama dengan tenor {duration}.')
    balance, interest_balance, previous_date = principal, total_interest, None
    output, zero_first = [], False
    for index, row in enumerate(rows, 1):
        context = f"Angsuran baris {row['_row']} ({header['Import Key']})"
        if number(row['Installment No.'], f'{context} No.', integer=True, positive=True) != index:
            raise ScheduleImportError(f'{context}: urutkan nomor angsuran 1 sampai {duration}, tanpa duplikasi/lompatan.')
        when = input_date(row['Due Date'], f'{context} Due Date')
        if (index == 1 and when != due) or (previous_date and when <= previous_date):
            raise ScheduleImportError(f'{context}: tanggal pertama harus sama dengan header, lalu meningkat tanpa duplikasi.')
        p = number(row['Principal'], f'{context} Principal', quantum=quantum)
        interest = number(row['Interest'], f'{context} Interest', quantum=quantum)
        total = number(row['Total'], f'{context} Total', quantum=quantum) if row.get('Total') not in (None, '') else None
        if index == 1 and header['Compute Method'] == 'Bunga Flat' and p == interest == 0 and total:
            if total != installment:
                raise ScheduleImportError(f'{context}: Total angsuran pertama Flat harus sama dengan Monthly Installment.')
            p, zero_first = total, True
        if total is not None and total != p + interest:
            raise ScheduleImportError(f'{context}: Total harus sama dengan Principal + Interest.')
        balance -= p
        interest_balance -= interest
        if balance < 0 or interest_balance < 0:
            raise ScheduleImportError(f'{context}: saldo menjadi negatif; periksa Total Hutang/Total Bunga dan nominal angsuran.')
        for column, expected in [('Principal Balance', balance), ('Interest Balance', interest_balance)]:
            if row.get(column) not in (None, '') and number(row[column], f'{context} {column}', quantum=quantum) != expected:
                raise ScheduleImportError(f'{context}: {column} tidak cocok dengan saldo berjalan ({expected}).')
        output.append({'date': when, 'principal': float(p), 'interest': float(interest), 'date_is_manual': True})
        previous_date = when
    if balance or interest_balance:
        raise ScheduleImportError(f'{label}: total pokok/bunga tidak sama dengan header; saldo akhir harus 0 (pokok {balance}, bunga {interest_balance}).')
    return {'rows': output, 'principal': float(principal), 'interest': float(total_interest),
            'installment': float(installment), 'rate': float(rate), 'duration': duration,
            'due': due, 'start': start, 'zero_first': zero_first}
