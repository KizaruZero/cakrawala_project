import base64
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from odoo.exceptions import AccessError, UserError
from odoo.tests import Form, tagged
from odoo.tests.common import new_test_user
from odoo.tools import file_path

from ..lib.schedule_import import HEADER_COLUMNS, LINE_COLUMNS, ScheduleImportError, read_workbook
from .test_early_settlement import LeasingScheduleCommon


def fixture_xlsx(header_rows, line_rows, formula=False):
    """Minimal OOXML fixture, not an alternative production template writer."""
    namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>''')
        archive.writestr('_rels/.rels', '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>''')
        archive.writestr('xl/workbook.xml', f'''<workbook xmlns="{namespace}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Header" sheetId="1" r:id="rId1"/><sheet name="Angsuran" sheetId="2" r:id="rId2"/></sheets></workbook>''')
        archive.writestr('xl/_rels/workbook.xml.rels', '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/></Relationships>''')
        for index, (columns, data) in enumerate([(HEADER_COLUMNS, header_rows), (LINE_COLUMNS, line_rows)], 1):
            root = ET.Element('worksheet', xmlns=namespace)
            sheet_data = ET.SubElement(root, 'sheetData')
            for number, values in enumerate([list(columns)] + [[record.get(column) for column in columns] for record in data], 1):
                row = ET.SubElement(sheet_data, 'row', r=str(number))
                for column_index, value in enumerate(values, 1):
                    if value is None:
                        continue
                    cell = ET.SubElement(row, 'c', r=f'{get_column_letter(column_index)}{number}')
                    if formula and index == 2 and number == 2 and column_index == 4:
                        ET.SubElement(cell, 'f').text = '200+300'
                    elif isinstance(value, (int, float)):
                        ET.SubElement(cell, 'v').text = str(value)
                    else:
                        cell.set('t', 'inlineStr')
                        ET.SubElement(ET.SubElement(cell, 'is'), 't').text = str(value)
            archive.writestr(f'xl/worksheets/sheet{index}.xml', ET.tostring(root))
    return output.getvalue()


@tagged('post_install', '-at_install')
class TestLeasingScheduleImport(LeasingScheduleCommon):
    def _data(self, key='IMP-EF-001', method='Bunga Efektif'):
        header = dict(zip(HEADER_COLUMNS, [
            key, method, f'AG-{key}', None, self.vendor.name, self.vehicle.license_plate,
            '2026-11-01', '2026-11-30', 3, 1000, 1500, 500, 400, 12, 100, 'Arrears',
            self.principal_account.code, self.accrued_account.code, self.expense_account.code,
            self.accrued_account.code, self.general_journal.code,
        ]))
        rows = [dict(zip(LINE_COLUMNS, row)) for row in [
            [key, 1, '2026-11-30', 300, 50, 350, 700, 50],
            [key, 2, '2026-12-30', 400, 30, 430, 300, 20],
            [key, 3, '2027-01-30', 300, 20, 320, 0, 0],
        ]]
        if method == 'Bunga Flat':
            header.update({'Monthly Installment': 100, 'Advarr': 'ADDM'})
            rows = [dict(zip(LINE_COLUMNS, row)) for row in [
                [key, 1, '2026-11-30', 0, 0, 100, 900, 100],
                [key, 2, '2026-12-30', 450, 50, 500, 450, 50],
                [key, 3, '2027-01-30', 450, 50, 500, 0, 0],
            ]]
        return [header], rows

    def _wizard(self, headers, rows, **extra):
        return self.env['leasing.schedule.import'].create({
            'file': base64.b64encode(fixture_xlsx(headers, rows)),
            'filename': 'contracts.xlsx', 'company_id': self.company.id, **extra,
        })

    def _import(self, method='Bunga Efektif', key='IMP-EF-001'):
        action = self._wizard(*self._data(key, method)).action_import()
        return self.env['account.loan'].browse(action['domain'][0][2])

    def test_all_methods_import_exact_contract_values_not_compute_formula(self):
        for method in ['Bunga Efektif', 'Bunga Efektif (Tanpa DP)', 'Bunga Flat']:
            loan = self._import(method, key=method)
            self.assertEqual(loan.state, 'draft')
            self.assertEqual(loan.schedule_source, 'import')
            self.assertEqual(loan.total_bunga, 100)
            self.assertEqual(loan.total_installment_payment, 1100)
            self.assertEqual(loan.saldo_pokok_hutang, 1000)
            self.assertEqual(sum(loan.line_ids.mapped('principal')), 1000)
            self.assertEqual(loan.line_ids[-1].outstanding_balance, 0)
            self.assertEqual(loan.line_ids[-1].interest_balance, 0)
            self.assertEqual(loan.agreement_no, f'AG-{method}')
            self.assertEqual(loan.vendor_id, self.vendor)
            self.assertFalse(loan.purchase_order_id)
            self.assertRegex(loan.name, r'^TA/CRS/\d{2}/\d{4}/\d{5}$')
            self.assertFalse(loan.line_ids.vendor_bill_id | loan.line_ids.monthly_adjustment_move_id)
            before = loan.line_ids.read(['date', 'principal', 'interest'])
            loan.action_confirm()
            self.assertEqual(loan.state, 'running')
            self.assertEqual(loan.line_ids.read(['date', 'principal', 'interest']), before)

    def test_flat_first_zero_display_is_real_principal_payment(self):
        loan = self._import('Bunga Flat')
        first = loan.line_ids[0]
        self.assertEqual((first.flat_display_principal, first.interest, first.payment), (0, 0, 100))
        self.assertEqual(first.principal, 100)
        self.assertEqual(first.outstanding_balance, 900)
        self.assertEqual(first.interest_balance, 100)
        loan.action_confirm()
        self.env['leasing.payment.wizard'].with_context(active_ids=first.ids).create({
            'journal_id': self.bank_journal.id, 'payment_date': date(2026, 11, 30),
        }).action_confirm_payment()
        self.assertEqual(first.payment_date, date(2026, 11, 30))
        self.assertEqual(loan.outstanding_balance, 900)
        self.assertEqual(self._bill(loan).amount_total, 900)

    def test_imported_effective_interest_posts_without_reallocation(self):
        loan = self._import()
        loan.action_confirm()
        first = loan.line_ids[0]
        first._generate_monthly_adjustment_entry()
        self.assertEqual(first.monthly_adjustment_move_id.amount_total, 50)
        self.assertFalse(first.payment_date)
        self.env['leasing.payment.wizard'].with_context(active_ids=first.ids).create({
            'journal_id': self.bank_journal.id, 'payment_date': date(2026, 11, 30),
        }).action_confirm_payment()
        self.assertEqual(first.vendor_bill_id.amount_total, 350)
        self.assertEqual(loan.outstanding_balance, 700)
        bill = self._bill(loan)
        self.assertEqual(bill.amount_total, 700)
        bill.action_post()

    def test_multiple_contracts_one_upload_and_duplicate_rejected_without_replace(self):
        h1, r1 = self._data()
        h2, r2 = self._data('IMP-FLAT-001', 'Bunga Flat')
        wizard = self._wizard(h1 + h2, r1 + r2)
        action = wizard.action_import()
        loans = self.env['account.loan'].browse(action['domain'][0][2])
        self.assertEqual(len(loans), 2)
        before = loans.line_ids.read(['date', 'principal', 'interest'])
        with self.assertRaisesRegex(UserError, 'sudah pernah'):
            wizard.action_import()
        self.assertEqual(loans.line_ids.read(['date', 'principal', 'interest']), before)

    def test_duplicate_existing_agreement_vendor_and_vehicle_rejected(self):
        self._import()
        headers, rows = self._data('OTHER-KEY')
        headers[0]['Agreement No.'] = 'AG-IMP-EF-001'
        with self.assertRaisesRegex(UserError, 'Agreement No.'):
            self._wizard(headers, rows).action_import()

    def test_archived_import_key_is_not_reused(self):
        self._import().active = False
        with self.assertRaisesRegex(UserError, 'sudah pernah'):
            self._wizard(*self._data()).action_import()

    def test_invalid_contracts_are_atomic_and_leave_no_new_loans(self):
        h1, r1 = self._data()
        h2, r2 = self._data('IMP-INVALID')
        h2[0]['Vendor'] = 'NONEXISTENT IMPORT VENDOR'
        before = self.env['account.loan'].search_count([])
        with self.assertRaisesRegex(UserError, 'Vendor'):
            self._wizard(h1 + h2, r1 + r2).action_import()
        self.assertEqual(self.env['account.loan'].search_count([]), before)

    def test_invalid_number_total_balance_date_and_duration_are_rejected(self):
        cases = [('Total Hutang', 1001), ('Total Bunga', 101), ('Duration (Months)', 4),
                 ('Monthly Installment', 0), ('Interest (%/Year)', -1), ('Start Date', '2026-12-01')]
        for field, value in cases:
            headers, rows = self._data()
            headers[0][field] = value
            with self.subTest(field=field), self.assertRaises(UserError):
                self._wizard(headers, rows).action_import()
        for field, value in [('Principal', -1), ('Interest', 0.001), ('Total', 1),
                             ('Principal Balance', 1), ('Interest Balance', 1),
                             ('Due Date', '11/30/2026'), ('Installment No.', 2)]:
            headers, rows = self._data()
            rows[0][field] = value
            with self.subTest(field=field), self.assertRaises(UserError):
                self._wizard(headers, rows).action_import()

    def test_duplicate_dates_unknown_keys_and_missing_rows_are_rejected(self):
        for mutation in ['date', 'key', 'missing']:
            headers, rows = self._data()
            if mutation == 'date':
                rows[1]['Due Date'] = rows[0]['Due Date']
            elif mutation == 'key':
                rows.append({**rows[-1], 'Import Key': 'NO-HEADER'})
            else:
                rows.pop()
            with self.subTest(mutation=mutation), self.assertRaises(UserError):
                self._wizard(headers, rows).action_import()

    def test_optional_totals_and_balances_are_derived_not_required(self):
        headers, rows = self._data()
        for row in rows:
            for key in ['Total', 'Principal Balance', 'Interest Balance']:
                row[key] = None
        action = self._wizard(headers, rows).action_import()
        self.assertEqual(self.env['account.loan'].browse(action['domain'][0][2]).total_installment_payment, 1100)

    def test_formula_invalid_binary_extension_and_empty_template_are_rejected(self):
        h, r = self._data()
        wizard = self._wizard(h, r)
        wizard.file = base64.b64encode(fixture_xlsx(h, r, formula=True))
        with self.assertRaisesRegex(UserError, 'formula'):
            wizard.action_import()
        for values in [{'file': base64.b64encode(b'not an xlsx')}, {'filename': 'file.xls'}]:
            wizard.write(values)
            with self.assertRaises(UserError):
                wizard.action_import()
        raw = Path(file_path('x_leasing_amortization/static/templates/Template_Import_Leasing.xlsx')).read_bytes()
        with self.assertRaisesRegex(ScheduleImportError, 'belum berisi'):
            read_workbook(raw)

    def test_ambiguous_vendor_and_missing_analytic_account_are_rejected(self):
        self.env['res.partner'].create({'name': self.vendor.name})
        with self.assertRaisesRegex(UserError, 'tidak unik'):
            self._wizard(*self._data()).action_import()
        self.env['res.partner'].search([('name', '=', self.vendor.name), ('id', '!=', self.vendor.id)]).unlink()
        self.vehicle.analytic_account_id = False
        with self.assertRaisesRegex(UserError, 'Analytic'):
            self._wizard(*self._data()).action_import()

    def test_wrong_account_type_or_journal_are_rejected(self):
        for field, value in [('Principal Account', self.payable_account.code),
                             ('Interest Expense Account', self.principal_account.code),
                             ('Interest Journal', self.bank_journal.code)]:
            h, r = self._data()
            h[0][field] = value
            with self.subTest(field=field), self.assertRaises(UserError):
                self._wizard(h, r).action_import()

    def test_imported_schedule_cannot_be_overwritten_by_compute_reset_or_copy(self):
        loan = self._import()
        before = self._snapshot(loan)
        for action in [loan.action_compute_amortization_schedule, loan.action_reset, loan.copy]:
            with self.assertRaises(UserError):
                action()
        self.assertEqual(self._snapshot(loan), before)

    def test_stale_header_or_changed_financial_rows_cannot_be_confirmed(self):
        loan = self._import()
        with self.assertRaisesRegex(UserError, 'file kontrak'):
            loan.line_ids[0].interest += 1
        loan.interest_rate_annual = 15
        with self.assertRaisesRegex(UserError, 'file kontrak'):
            loan.action_confirm()
        with self.assertRaisesRegex(UserError, 'file kontrak'):
            loan.line_ids[0].date = date(2026, 11, 29)

    def test_native_draft_date_edit_preserves_imported_amounts(self):
        loan = self._import('Bunga Flat')
        before = loan.line_ids.read(['principal', 'interest'])
        with Form(loan, view='account_loans.account_loan_form_view') as form:
            with form.flat_line_ids.edit(1) as row:
                row.date = date(2026, 12, 31)
        self.assertEqual(loan.line_ids.read(['principal', 'interest']), before)
        loan.action_confirm()

    def test_first_manual_due_date_updates_import_header_without_changing_amounts(self):
        loan = self._import()
        before = loan.line_ids.read(['principal', 'interest'])
        with Form(loan, view='account_loans.account_loan_form_view') as form:
            with form.line_ids.edit(0) as row:
                row.date = date(2026, 11, 29)
        self.assertEqual(loan.due_date, date(2026, 11, 29))
        self.assertEqual(loan.line_ids.read(['principal', 'interest']), before)
        loan.action_confirm()

    def test_running_imported_rows_and_header_protected_and_payment_draft_blocked(self):
        loan = self._import()
        wizard = self.env['leasing.payment.wizard'].with_context(active_ids=loan.line_ids[:1].ids).create({
            'journal_id': self.bank_journal.id})
        with self.assertRaisesRegex(UserError, 'Running'):
            wizard.action_confirm_payment()
        loan.action_confirm()
        for action in [lambda: loan.write({'duration': 4}), lambda: loan.line_ids[0].write({'principal': 5}),
                       loan.line_ids[0].unlink,
                       lambda: self.env['account.loan.line'].create({'loan_id': loan.id, 'date': date(2030, 1, 1)})]:
            with self.assertRaises(UserError):
                action()

    def test_import_metadata_cannot_be_spoofed_by_rpc_context(self):
        with self.assertRaises(UserError):
            self.env['account.loan'].with_context(_leasing_import_token=True).create({'schedule_source': 'import'})
        loan = self._import()
        for values in [{'schedule_source': 'compute'}, {'import_key': 'NEW'}, {'import_schedule_signature': 'forged'}]:
            with self.assertRaises(UserError):
                loan.write(values)

    def test_fleet_only_user_and_inactive_company_cannot_import(self):
        wizard = self._wizard(*self._data())
        fleet = new_test_user(self.env, login='schedule_import_fleet', groups='fleet.fleet_group_user')
        with self.assertRaises(AccessError):
            wizard.with_user(fleet).action_import()
        other = self.env['res.company'].create({'name': 'Inactive Import Company'})
        wizard.company_id = other
        with self.assertRaises(AccessError):
            wizard.with_context(allowed_company_ids=self.company.ids).action_import()

    def test_effective2_import_payment_and_settlement_use_file_not_formula(self):
        loan = self._import('Bunga Efektif (Tanpa DP)')
        loan.action_confirm()
        first = loan.line_ids[0]
        first._generate_monthly_adjustment_entry()
        self.assertEqual(first.monthly_adjustment_move_id.amount_total, 50)
        self.env['leasing.payment.wizard'].with_context(active_ids=first.ids).create({
            'journal_id': self.bank_journal.id, 'payment_date': date(2026, 11, 30),
        }).action_confirm_payment()
        self.assertEqual(first.vendor_bill_id.amount_total, 350)
        bill = self._bill(loan)
        self.assertEqual(bill.amount_total, 700)
        bill.action_post()

    def test_legacy_effective2_template_label_remains_supported(self):
        loan = self._import('Bunga Efektif 2')
        self.assertEqual(loan.compute_method, 'effective_2')
        self.assertEqual(dict(loan._fields['compute_method'].selection)['effective_2'],
                         'Bunga Efektif (Tanpa DP)')
        self.assertEqual(loan.total_installment_payment, 1100)
        self.assertEqual(loan.line_ids.mapped('principal'), [300, 400, 300])
        loan.action_confirm()

    def test_import_rejects_duplicate_key_within_workbook_and_wrong_company_vehicle(self):
        headers, rows = self._data()
        with self.assertRaisesRegex(UserError, 'unik'):
            self._wizard(headers + headers, rows).action_import()
        other = self.env['res.company'].create({'name': 'Foreign Vehicle Company'})
        self.vehicle.company_id = other
        with self.assertRaisesRegex(UserError, 'Plate Number'):
            self._wizard(headers, rows).action_import()

    def test_template_matches_parser_and_contains_validations_not_sample_contracts(self):
        path = file_path('x_leasing_amortization/static/templates/Template_Import_Leasing.xlsx')
        workbook = load_workbook(path)
        self.assertEqual(tuple(cell.value for cell in workbook['Header'][1]), HEADER_COLUMNS)
        self.assertEqual(tuple(cell.value for cell in workbook['Angsuran'][1]), LINE_COLUMNS)
        self.assertTrue(all(cell.value is None for row in workbook['Header'].iter_rows(min_row=2) for cell in row))
        self.assertEqual(len(workbook['Header'].data_validations.dataValidation), 2)
        method_rule = next(rule for rule in workbook['Header'].data_validations.dataValidation
                           if str(rule.sqref) == 'B2:B21')
        self.assertIn('Bunga Efektif (Tanpa DP)', method_rule.formula1)
        self.assertNotIn('Bunga Efektif 2', method_rule.formula1)
        self.assertEqual(workbook['Header'].freeze_panes, 'B2')
        workbook.close()
        arch = etree.fromstring(self.env['leasing.schedule.import'].get_view(view_type='form')['arch'])
        self.assertTrue(arch.xpath('//button[@name="action_download_template"]'))
